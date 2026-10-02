from openpyxl import load_workbook

# Open the workbook
workbook = load_workbook("centrage.xlsx")

# Print all sheet names
print(workbook.sheetnames)